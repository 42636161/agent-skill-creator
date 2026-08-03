#!/usr/bin/env python3
"""
Generate realistic Chinese retail enterprise test data for the agent-skill-creator
performance test suite. Each function corresponds to one test case (A1-E12).

Usage:
    python3 tests/retail_skills/generate_data.py          # generate all
    python3 tests/retail_skills/generate_data.py --case a1 # single case
"""

import csv
import os
import random
import sqlite3
import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

DATA_DIR = Path(__file__).resolve().parent / "data"
random.seed(42)

# ── helpers ────────────────────────────────────────────────────────────────

STORES = ["朝阳大悦城店", "西单店", "三里屯店", "五棵松店", "望京店", "通州万达店"]
CATEGORIES = {
    "生鲜": ["蔬菜", "水果", "肉禽", "水产", "蛋奶"],
    "食品": ["零食", "饮料", "粮油", "调味品", "冷冻食品"],
    "日用品": ["洗护", "纸品", "清洁", "家杂", "一次性用品"],
    "纺织": ["床品", "毛巾", "拖鞋", "家居服", "窗帘"],
}
SKU_NAMES = {
    "蔬菜": ["有机菠菜", "西红柿500g", "黄瓜", "生菜", "青椒"],
    "水果": ["富士苹果", "进口香蕉", "巨峰葡萄", "海南芒果", "麒麟西瓜"],
    "肉禽": ["冷鲜鸡胸肉", "五花肉500g", "牛腱子", "羊排", "鸭腿"],
    "水产": ["活鲫鱼", "基围虾", "三文鱼切片", "带鱼段", "花蛤"],
    "蛋奶": ["鲜牛奶1L", "酸奶8连杯", "土鸡蛋10枚", "鹌鹑蛋", "奶酪片"],
    "零食": ["薯片大包", "坚果混合装", "牛肉干", "巧克力", "夹心饼干"],
    "饮料": ["可乐330ml", "无糖茶饮", "椰子水", "功能饮料", "矿泉水"],
    "粮油": ["五常大米5kg", "橄榄油500ml", "挂面1kg", "小米", "糯米"],
    "调味品": ["生抽500ml", "蚝油", "料酒", "豆瓣酱", "花椒油"],
    "冷冻食品": ["速冻水饺", "汤圆", "手抓饼", "冰淇淋", "冷冻虾仁"],
    "洗护": ["洗发水500ml", "沐浴露", "护发素", "洗手液", "牙膏"],
    "纸品": ["抽纸3连包", "卷纸12卷", "湿巾", "厨房纸", "手帕纸"],
    "清洁": ["洗洁精", "洗衣液", "地板清洁剂", "马桶清洁剂", "玻璃水"],
    "家杂": ["保鲜膜", "垃圾袋", "密封袋", "一次性手套", "铝箔"],
    "一次性用品": ["纸杯", "一次性筷子", "吸管", "餐盒", "手套"],
    "床品": ["四件套", "夏凉被", "枕头", "床笠", "毛毯"],
    "毛巾": ["面巾", "浴巾", "方巾套装", "干发帽", "运动毛巾"],
    "拖鞋": ["棉拖鞋", "凉拖鞋", "浴室拖鞋", "家居拖鞋", "儿童拖鞋"],
    "家居服": ["女士家居服", "男士家居服", "儿童睡衣", "情侣套装", "浴袍"],
    "窗帘": ["遮光窗帘", "纱帘", "百叶窗", "卷帘", "罗马帘"],
}

def _sku_id(cat: str, sub: str, idx: int) -> str:
    """Generate a realistic SKU ID like SP001234."""
    prefix = {"生鲜":"SX","食品":"SP","日用品":"RY","纺织":"FZ"}[cat]
    return f"{prefix}{idx:06d}"

def _price(sub_cat: str) -> float:
    """Realistic retail prices in CNY."""
    prices = {"蔬菜": (3,12), "水果": (5,30), "肉禽": (12,45), "水产": (8,60),
              "蛋奶": (6,25), "零食": (4,20), "饮料": (2,8), "粮油": (15,60),
              "调味品": (5,18), "冷冻食品": (8,35), "洗护": (15,50),
              "纸品": (8,30), "清洁": (6,25), "家杂": (3,12),
              "一次性用品": (2,8), "床品": (80,400), "毛巾": (15,80),
              "拖鞋": (10,40), "家居服": (60,200), "窗帘": (50,300)}
    lo, hi = prices.get(sub_cat, (5,30))
    return round(random.uniform(lo, hi), 2)

def _cost(price: float) -> float:
    return round(price * random.uniform(0.50, 0.75), 2)

# ── A1: 门店日报 ───────────────────────────────────────────────────────────

def generate_a1():
    """Multi-sheet Excel: 销售明细, 退货, 会员, 促销."""
    wb = openpyxl.Workbook()
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"))

    def style_sheet(ws, headers, rows):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(1, c, h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")
        for r, row in enumerate(rows, 2):
            for c, val in enumerate(row, 1):
                cell = ws.cell(r, c, val)
                cell.border = thin_border

    # Sheet 1: 销售明细
    ws1 = wb.active
    ws1.title = "销售明细"
    sales_headers = ["日期","门店","SKU编号","商品名称","品类","子品类","销售数量","单价","销售金额","支付方式","会员ID"]
    sales_rows = []
    for day_offset in range(31):
        d = date(2026, 7, 1) + timedelta(days=day_offset)
        n_txns = random.randint(40, 80)
        for _ in range(n_txns):
            store = random.choice(STORES)
            cat = random.choice(list(CATEGORIES))
            sub = random.choice(CATEGORIES[cat])
            sku_names = SKU_NAMES[sub]
            name = random.choice(sku_names)
            qty = random.randint(1, 15)
            price = _price(sub)
            amt = round(qty * price, 2)
            method = random.choice(["微信","支付宝","银行卡","现金","会员卡"])
            member = f"M{random.randint(10000,99999)}" if random.random()<0.6 else ""
            sales_rows.append([str(d), store, _sku_id(cat, sub, sku_names.index(name)),
                              name, cat, sub, qty, price, amt, method, member])
    style_sheet(ws1, sales_headers, sales_rows)

    # Sheet 2: 退货
    ws2 = wb.create_sheet("退货")
    ret_headers = ["退货日期","门店","SKU编号","商品名称","退货数量","退货金额","退货原因","原销售日期"]
    reasons = ["质量问题","过期","包装破损","顾客不满意","错发漏发","其他"]
    ret_rows = []
    for day_offset in range(31):
        d = date(2026, 7, 1) + timedelta(days=day_offset)
        for _ in range(random.randint(2, 8)):
            store = random.choice(STORES)
            cat = random.choice(list(CATEGORIES))
            sub = random.choice(CATEGORIES[cat])
            sku_names = SKU_NAMES[sub]
            name = random.choice(sku_names)
            qty = random.randint(1, 3)
            price = _price(sub)
            orig_date = d - timedelta(days=random.randint(1, 7))
            ret_rows.append([str(d), store, _sku_id(cat, sub, sku_names.index(name)),
                            name, qty, round(qty*price,2), random.choice(reasons),
                            str(orig_date)])
    style_sheet(ws2, ret_headers, ret_rows)

    # Sheet 3: 会员
    ws3 = wb.create_sheet("会员")
    mem_headers = ["会员ID","姓名","性别","手机号","注册日期","会员等级","累计积分","本月消费金额"]
    levels = ["普通","银卡","金卡","钻石"]
    mem_rows = []
    for i in range(200):
        mid = f"M{10000+i}"
        name = f"顾客{random.choice('赵钱孙李周吴郑王冯陈')}{random.choice('伟芳娜敏强磊洋')}"
        reg_d = date(2025, random.randint(1,12), random.randint(1,28))
        lv = random.choices(levels, weights=[40,35,20,5])[0]
        pts = random.randint(0, 50000)
        monthly = round(random.uniform(50, 5000), 2) if lv != "普通" else round(random.uniform(0,300),2)
        mem_rows.append([mid, name, random.choice(["男","女"]),
                        f"1{random.randint(30,99)}{random.randint(10000000,99999999)}",
                        str(reg_d), lv, pts, monthly])
    style_sheet(ws3, mem_headers, mem_rows)

    # Sheet 4: 促销
    ws4 = wb.create_sheet("促销")
    promo_headers = ["促销活动名称","开始日期","结束日期","促销类型","适用品类","折扣率","参与SKU数","预算金额"]
    promo_types = ["满减","折扣","买赠","会员专享","限时秒杀","第二件半价"]
    promo_data = [
        ["暑期清凉节", "2026-07-01","2026-07-15","满减","饮料/冷冻食品",0.15,45,50000],
        ["会员日特惠", "2026-07-08","2026-07-08","会员专享","全品类",0.20,120,30000],
        ["周末生鲜购", "2026-07-11","2026-07-12","折扣","生鲜",0.12,30,15000],
        ["年中床品大促", "2026-07-15","2026-07-31","满减","纺织",0.25,25,80000],
        ["零食节", "2026-07-20","2026-07-27","第二件半价","零食",0.30,60,20000],
    ]
    for row in promo_data:
        promo_rows = promo_data
    style_sheet(ws4, promo_headers, promo_data)

    out = DATA_DIR / "a1_daily_report" / "门店销售_202607.xlsx"
    wb.save(out)
    print(f"  ✓ A1: {out} ({len(sales_rows)} sales rows, {len(ret_rows)} returns, {len(mem_rows)} members)")


# ── A3: ERP 对账 ────────────────────────────────────────────────────────────

def generate_a3():
    """Two CSV files: ERP export and POS data with intentional mismatches."""
    def write_csv(name, headers, rows):
        p = DATA_DIR / "a3_erp_recon" / name
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(headers)
            w.writerows(rows)
        return p

    # 金蝶导出
    erp_headers = ["单据日期","单据编号","门店","商品编码","商品名称","销售数量","销售金额","收款方式","经手人"]
    erp_rows = []
    for day_offset in range(7):
        d = date(2026, 7, 20) + timedelta(days=day_offset)
        for i in range(random.randint(30, 50)):
            store = random.choice(STORES)
            cat = random.choice(list(CATEGORIES))
            sub = random.choice(CATEGORIES[cat])
            name = random.choice(SKU_NAMES[sub])
            qty = random.randint(1, 10)
            amt = round(qty * _price(sub), 2)
            erp_rows.append([str(d), f"JD-{d.strftime('%Y%m%d')}-{i:04d}", store,
                            _sku_id(cat, sub, SKU_NAMES[sub].index(name)),
                            name, qty, amt, random.choice(["现金","微信","银行卡"]),
                            random.choice(["张丽","王强","李芳"])])
    p1 = write_csv("金蝶_导出_202607.csv", erp_headers, erp_rows)

    # POS data - mostly matching, some differences
    pos_headers = ["交易时间","交易号","店铺","商品条码","品名","数量","实收金额","支付方式","收银员"]
    pos_rows = []
    # copy 90% from ERP with slight variations, add 10% unique
    for row in erp_rows[:int(len(erp_rows)*0.9)]:
        pos_rows.append([row[0]+" 14:30:00", row[1], row[2], row[3],
                        row[4], row[5], round(row[6]*random.uniform(0.98,1.0),2),
                        row[7], random.choice(["小刘","小陈","小赵"])])
    # 10% POS-only transactions
    for _ in range(int(len(erp_rows)*0.1)):
        d = date(2026, 7, 20) + timedelta(days=random.randint(0,6))
        store = random.choice(STORES)
        sub = random.choice(CATEGORIES["食品"])
        name = random.choice(SKU_NAMES[sub])
        qty = random.randint(1,5)
        pos_rows.append([str(d)+" 18:00:00", f"POS-{random.randint(10000,99999)}",
                        store, _sku_id("食品", sub, 0), name, qty,
                        round(qty*_price(sub),2), "微信", "小刘"])
    p2 = write_csv("pos_sales_202607.csv", pos_headers, pos_rows)
    print(f"  ✓ A3: {p1} ({len(erp_rows)} rows), {p2} ({len(pos_rows)} rows)")


# ── B4: 促销 ROI ────────────────────────────────────────────────────────────

def generate_b4():
    p = DATA_DIR / "b4_promo_roi" / "promo_events.csv"
    headers = ["活动编号","活动名称","开始日期","结束日期","覆盖门店","适用品类","促销类型",
               "折扣力度","营销费用","活动前日均销售额","活动中日均销售额","活动后一周日均销售额"]
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for i in range(1, 16):
            start = date(2026, 6, 1) + timedelta(days=random.randint(0,50))
            dur = random.choice([1,3,5,7])
            end = start + timedelta(days=dur)
            base_sales = random.uniform(3000, 15000)
            lift = random.uniform(0.1, 0.8)  # promotion lift
            post = random.uniform(-0.1, 0.3)  # post-promo effect
            w.writerow([
                f"P2026{i:04d}",
                random.choice(["周末大促","品牌日","品类节","新店开业","店庆回馈","清仓特卖",
                              "会员专享日","双十一预热","年中年终大促"]),
                str(start), str(end),
                ",".join(random.sample(STORES, k=random.randint(1,3))),
                random.choice(["全品类","生鲜","食品","日用品","纺织"]),
                random.choice(["满减","折扣","买赠","秒杀","第二件半价"]),
                round(random.uniform(0.10, 0.40), 2),
                round(random.uniform(2000, 30000), 2),
                round(base_sales, 2),
                round(base_sales*(1+lift), 2),
                round(base_sales*(1+post), 2),
            ])
    print(f"  ✓ B4: {p} (15 promo events)")


# ── B5: 商品汰换 ──────────────────────────────────────────────────────────

def generate_b5():
    p = DATA_DIR / "b5_assortment" / "product_performance.csv"
    headers = ["SKU编号","商品名称","品类","子品类","供应商","进货价","零售价","季度销量",
               "季度销售额","毛利率","库存周转天数","退货率","是否季节性商品","上架日期"]
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        idx = 0
        for cat, subs in CATEGORIES.items():
            for sub in subs:
                for name in SKU_NAMES[sub]:
                    idx += 1
                    price = _price(sub)
                    cost = _cost(price)
                    qty = random.randint(20, 500)
                    margin = round((price-cost)/price, 4)
                    turn_days = round(random.uniform(5, 120) / (1+margin*3), 1)
                    ret_rate = round(random.uniform(0.001, 0.08) * (1-margin), 4)
                    w.writerow([
                        _sku_id(cat, sub, SKU_NAMES[sub].index(name)),
                        name, cat, sub,
                        f"供应商{random.choice('ABCDEFGH')}",
                        cost, price, qty, round(price*qty, 2),
                        margin, turn_days, ret_rate,
                        "是" if random.random()<0.2 else "否",
                        str(date(2025, random.randint(1,12), random.randint(1,28)))
                    ])
    print(f"  ✓ B5: {p} ({idx} SKUs)")


# ── B6: 会员 RFM ───────────────────────────────────────────────────────────

def generate_b6():
    p = DATA_DIR / "b6_rfm" / "transactions.csv"
    headers = ["交易ID","会员ID","交易日期","交易金额","交易门店","购买品类数"]
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        member_ids = [f"M{10000+i}" for i in range(500)]
        for i in range(3000):
            mid = random.choice(member_ids)
            d = date(2026, 1, 1) + timedelta(days=random.randint(0, 210))
            amt = round(random.expovariate(1/200) * 200 + 15, 2)
            w.writerow([
                f"T{i:06d}", mid, str(d), min(amt, 5000),
                random.choice(STORES), random.randint(1, 5)
            ])
    print(f"  ✓ B6: {p} (3000 transactions, 500 members)")


# ── C7: 月度经营分析 ───────────────────────────────────────────────────────

def generate_c7():
    # Sales target
    p1 = DATA_DIR / "c7_monthly" / "sales_target.csv"
    with open(p1, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["月份","门店","目标销售额","实际销售额","达成率"])
        for m in range(1,7):
            for store in STORES:
                target = round(random.uniform(200000, 800000), 2)
                actual = round(target * random.uniform(0.75, 1.25), 2)
                w.writerow([f"2026-{m:02d}", store, target, actual, round(actual/target, 4)])

    # Inventory health
    p2 = DATA_DIR / "c7_monthly" / "inventory_health.csv"
    with open(p2, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["月份","门店","SKU总数","动销SKU数","滞销SKU数","库存金额","库销比","缺货次数"])
        for m in range(1,7):
            for store in STORES:
                total = random.randint(500, 2000)
                active = random.randint(int(total*0.6), total)
                dead = total - active
                inv_val = round(random.uniform(100000, 800000), 2)
                w.writerow([f"2026-{m:02d}", store, total, active, dead, inv_val,
                           round(random.uniform(0.8, 2.5), 2), random.randint(0,30)])

    # Staff efficiency
    p3 = DATA_DIR / "c7_monthly" / "staff_efficiency.csv"
    with open(p3, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["月份","门店","员工总数","总工时","人均销售额","人均服务客数","离职人数"])
        for m in range(1,7):
            for store in STORES:
                staff = random.randint(8, 40)
                hours = staff * random.randint(160, 200)
                w.writerow([f"2026-{m:02d}", store, staff, hours,
                           round(random.uniform(8000, 35000),2),
                           random.randint(200, 800),
                           random.randint(0,3)])

    # Cost control
    p4 = DATA_DIR / "c7_monthly" / "cost_control.csv"
    with open(p4, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["月份","门店","人力成本","租金","水电费","营销费用","损耗金额","总费用","费用率"])
        for m in range(1,7):
            for store in STORES:
                labor = round(random.uniform(30000, 150000), 2)
                rent = round(random.uniform(20000, 100000), 2)
                util = round(random.uniform(3000, 15000), 2)
                mktg = round(random.uniform(1000, 20000), 2)
                loss = round(random.uniform(500, 8000), 2)
                total = labor + rent + util + mktg + loss
                w.writerow([f"2026-{m:02d}", store, labor, rent, util, mktg, loss, total,
                           round(random.uniform(0.15, 0.35), 4)])
    print(f"  ✓ C7: 4 files in c7_monthly/")


# ── C8: 单店排班 ────────────────────────────────────────────────────────────

def generate_c8():
    p = DATA_DIR / "c8_scheduling" / "foot_traffic.csv"
    headers = ["日期","时间段","客流量","是否为节假日","天气"]
    weathers = ["晴","多云","小雨","大雨","阴"]
    holidays = {"2026-07-04":"周六","2026-07-05":"周日","2026-07-11":"周六",
                "2026-07-12":"周日","2026-07-18":"周六","2026-07-19":"周日",
                "2026-07-25":"周六","2026-07-26":"周日"}
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for day_offset in range(31):
            d = date(2026, 7, 1) + timedelta(days=day_offset)
            ds = str(d)
            is_holiday = "是" if ds in holidays else "否"
            base = 80 if is_holiday == "是" else 40
            for hour in [10,11,12,13,14,15,16,17,18,19,20,21]:
                # rush hour peaks
                peak = 1.0
                if hour in [11,12,18,19]:
                    peak = 2.5 if is_holiday=="是" else 1.8
                elif hour in [17,20]:
                    peak = 2.0 if is_holiday=="是" else 1.5
                traffic = int(base * peak * random.uniform(0.7, 1.3))
                w.writerow([ds, f"{hour}:00-{hour+1}:00", traffic, is_holiday, random.choice(weathers)])
    print(f"  ✓ C8: {p} (31 days × 12 hours)")


# ── D9: 动销率 ──────────────────────────────────────────────────────────────

def generate_d9():
    p = DATA_DIR / "d9_sell_through" / "sku_sales.csv"
    headers = ["SKU编号","商品名称","品类","子品类","门店","月初库存","月末库存","月销量","月销售额","上架天数"]
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        idx = 0
        for cat, subs in CATEGORIES.items():
            for sub in subs:
                for name in SKU_NAMES[sub]:
                    for store in random.sample(STORES, k=random.randint(2,5)):
                        idx += 1
                        begin = random.randint(10, 200)
                        sold = random.randint(0, int(begin*1.5))
                        end = max(0, begin - sold + random.randint(0, 20))
                        price = _price(sub)
                        w.writerow([
                            _sku_id(cat, sub, SKU_NAMES[sub].index(name)),
                            name, cat, sub, store,
                            begin, end, sold, round(sold*price, 2),
                            random.randint(1, 365)
                        ])
    print(f"  ✓ D9: {p} ({idx} SKU-store combinations)")


# ── E11: 薪资核算 ───────────────────────────────────────────────────────────

def generate_e11():
    # Employee roster
    p1 = DATA_DIR / "e11_payroll" / "employees.csv"
    roles = ["店长","副店长","收银员","理货员","生鲜技工","促销员","保洁"]
    base_pays = {"店长":8000,"副店长":6000,"收银员":4500,"理货员":4200,"生鲜技工":5000,"促销员":4000,"保洁":3500}
    with open(p1, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["员工编号","姓名","门店","岗位","基本工资","入职日期","社保基数"])
        for i in range(50):
            role = random.choice(roles)
            base = base_pays[role]
            w.writerow([f"E{i:04d}", f"员工{random.choice('赵钱孙李周吴郑王')}{random.choice('明华芳伟')}",
                       random.choice(STORES), role, base,
                       str(date(2020,1,1)+timedelta(days=random.randint(0,2000))),
                       round(base*random.uniform(0.8,1.2),2)])

    # Commission rules (Markdown)
    p2 = DATA_DIR / "e11_payroll" / "commission_rules.md"
    rules = """# 门店员工提成规则

## 适用岗位
店长、副店长、收银员、理货员、促销员

## 提成计算方式

### 店长
- 门店月销售额达成率 >= 100%：超出部分 × 0.5%
- 门店月销售额达成率 >= 120%：超出部分 × 0.8%（叠加）

### 副店长
- 门店月销售额达成率 >= 100%：超出部分 × 0.3%

### 收银员
- 个人经手交易额 × 0.1%

### 理货员
- 所负责品类销售额 × 0.15%

### 促销员
- 促销活动期间销售额 × 0.5%
- 非促销期间销售额 × 0.2%

## 全勤奖
当月无迟到、早退、请假：+300 元

## 加班费
- 工作日加班：基本工资 / 21.75 / 8 × 1.5
- 休息日加班：基本工资 / 21.75 / 8 × 2.0
- 法定节假日加班：基本工资 / 21.75 / 8 × 3.0

## 扣款规则
- 迟到一次：扣 20 元
- 早退一次：扣 30 元
- 旷工一天：扣 3 天工资
- 事假一天：当日基本工资不发放
- 病假一天：当日基本工资 × 60%

## 社保扣款
个人承担部分 = 社保基数 × (养老8% + 医疗2% + 失业0.5%)
"""
    with open(p2, "w", encoding="utf-8") as f:
        f.write(rules)
    print(f"  ✓ E11: {p1} (50 employees), {p2}")


# ── E12: 多格式周报 ─────────────────────────────────────────────────────────

def generate_e12():
    # Store sales
    p1 = DATA_DIR / "e12_multi_format" / "store_sales.csv"
    with open(p1, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["日期","门店","客流量","交易笔数","客单价","销售额","毛利率","同比销售额","环比销售额"])
        for day_offset in range(7):
            d = date(2026, 7, 24) + timedelta(days=day_offset)
            for store in STORES:
                traffic = random.randint(200, 1500)
                txns = random.randint(int(traffic*0.3), int(traffic*0.7))
                asp = round(random.uniform(35, 120), 2)
                sales = round(txns * asp, 2)
                margin = round(random.uniform(0.18, 0.42), 4)
                yoy = round(sales * random.uniform(0.85, 1.2), 2)
                mom = round(sales * random.uniform(0.9, 1.15), 2)
                w.writerow([str(d), store, traffic, txns, asp, sales, margin, yoy, mom])

    # Customer feedback
    p2 = DATA_DIR / "e12_multi_format" / "customer_feedback.csv"
    aspects = ["商品质量","价格","服务态度","环境卫生","排队时间","停车便利"]
    with open(p2, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["日期","门店","评价维度","评分","评价内容"])
        for day_offset in range(7):
            d = date(2026, 7, 24) + timedelta(days=day_offset)
            for store in STORES:
                for asp in random.sample(aspects, k=3):
                    score = random.randint(2, 5)
                    comments = {5:"非常满意",4:"比较满意",3:"一般",2:"不满意"}
                    w.writerow([str(d), store, asp, score, comments[score]])

    # Inventory snapshot
    p3 = DATA_DIR / "e12_multi_format" / "inventory_snapshot.csv"
    with open(p3, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["门店","品类","SKU数","库存数量","库存金额","低库存SKU数","过期预警SKU数"])
        for store in STORES:
            for cat in CATEGORIES:
                sku_count = random.randint(20, 150)
                qty = random.randint(200, 5000)
                val = round(random.uniform(5000, 200000), 2)
                low = random.randint(0, max(1, int(sku_count*0.15)))
                exp = random.randint(0, max(1, int(sku_count*0.05)))
                w.writerow([store, cat, sku_count, qty, val, low, exp])
    print(f"  ✓ E12: 3 files in e12_multi_format/")


# ── main ────────────────────────────────────────────────────────────────────

ALL_CASES = {
    "a1": generate_a1,
    "a3": generate_a3,
    "b4": generate_b4,
    "b5": generate_b5,
    "b6": generate_b6,
    "c7": generate_c7,
    "c8": generate_c8,
    "d9": generate_d9,
    "e11": generate_e11,
    "e12": generate_e12,
}

if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--case":
        case = sys.argv[2]
        if case in ALL_CASES:
            ALL_CASES[case]()
        else:
            print(f"Unknown case: {case}")
            sys.exit(1)
    else:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        for sub in ["a1_daily_report","a3_erp_recon","b4_promo_roi","b5_assortment",
                     "b6_rfm","c7_monthly","c8_scheduling","d9_sell_through",
                     "e11_payroll","e12_multi_format"]:
            (DATA_DIR/sub).mkdir(parents=True, exist_ok=True)
        print("Generating retail test data...\n")
        for name, fn in ALL_CASES.items():
            fn()
        print("\nAll test data generated.")
