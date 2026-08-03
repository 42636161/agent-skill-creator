#!/usr/bin/env python3
"""Generate HOLDOUT test data for evaluating generated skills — NEVER seen by creator."""
import csv, random
from datetime import date, timedelta
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DATA_DIR = Path(__file__).resolve().parent / "data_test"
random.seed(99)

STORES = ["朝阳大悦城店","西单店","三里屯店","五棵松店","望京店","通州万达店"]
CATEGORIES = {"生鲜":["蔬菜","水果","肉禽","水产","蛋奶"],"食品":["零食","饮料","粮油","调味品","冷冻食品"],"日用品":["洗护","纸品","清洁","家杂","一次性用品"],"纺织":["床品","毛巾","拖鞋","家居服","窗帘"]}
SKU_NAMES = {"蔬菜":["有机菠菜","西红柿500g","黄瓜","生菜","青椒"],"水果":["富士苹果","进口香蕉","巨峰葡萄","海南芒果","麒麟西瓜"],"肉禽":["冷鲜鸡胸肉","五花肉500g","牛腱子","羊排","鸭腿"],"水产":["活鲫鱼","基围虾","三文鱼切片","带鱼段","花蛤"],"蛋奶":["鲜牛奶1L","酸奶8连杯","土鸡蛋10枚","鹌鹑蛋","奶酪片"],"零食":["薯片大包","坚果混合装","牛肉干","巧克力","夹心饼干"],"饮料":["可乐330ml","无糖茶饮","椰子水","功能饮料","矿泉水"],"粮油":["五常大米5kg","橄榄油500ml","挂面1kg","小米","糯米"],"调味品":["生抽500ml","蚝油","料酒","豆瓣酱","花椒油"],"冷冻食品":["速冻水饺","汤圆","手抓饼","冰淇淋","冷冻虾仁"],"洗护":["洗发水500ml","沐浴露","护发素","洗手液","牙膏"],"纸品":["抽纸3连包","卷纸12卷","湿巾","厨房纸","手帕纸"],"清洁":["洗洁精","洗衣液","地板清洁剂","马桶清洁剂","玻璃水"],"家杂":["保鲜膜","垃圾袋","密封袋","一次性手套","铝箔"],"一次性用品":["纸杯","一次性筷子","吸管","餐盒","手套"],"床品":["四件套","夏凉被","枕头","床笠","毛毯"],"毛巾":["面巾","浴巾","方巾套装","干发帽","运动毛巾"],"拖鞋":["棉拖鞋","凉拖鞋","浴室拖鞋","家居拖鞋","儿童拖鞋"],"家居服":["女士家居服","男士家居服","儿童睡衣","情侣套装","浴袍"],"窗帘":["遮光窗帘","纱帘","百叶窗","卷帘","罗马帘"]}

def _sku_id(cat, sub, idx):
    return {"生鲜":"SX","食品":"SP","日用品":"RY","纺织":"FZ"}[cat] + f"{idx:06d}"

def _price(sub_cat):
    p = {"蔬菜":(3,12),"水果":(5,30),"肉禽":(12,45),"水产":(8,60),"蛋奶":(6,25),"零食":(4,20),"饮料":(2,8),"粮油":(15,60),"调味品":(5,18),"冷冻食品":(8,35),"洗护":(15,50),"纸品":(8,30),"清洁":(6,25),"家杂":(3,12),"一次性用品":(2,8),"床品":(80,400),"毛巾":(15,80),"拖鞋":(10,40),"家居服":(60,200),"窗帘":(50,300)}
    lo, hi = p.get(sub_cat, (5,30))
    return round(random.uniform(lo, hi), 2)

def write_csv(dn, fn, hdrs, rows):
    p = DATA_DIR / dn / fn
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(hdrs); w.writerows(rows)
    return p

# ── Excel helper ──
def xl_style(ws, hdrs, rows):
    hf = Font(bold=True, size=11); hfill = PatternFill("solid", fgColor="D9E1F2")
    bd = Border(left=Side(style="thin"),right=Side(style="thin"),top=Side(style="thin"),bottom=Side(style="thin"))
    for c, h in enumerate(hdrs, 1):
        cl = ws.cell(1, c, h); cl.font = hf; cl.fill = hfill; cl.border = bd; cl.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1): ws.cell(r, c, val).border = bd

# ════════════════════════════════════════════════════════════
# A1
# ════════════════════════════════════════════════════════════
def gen_a1():
    wb = openpyxl.Workbook()
    hd = ["日期","门店","SKU编号","商品名称","品类","子品类","销售数量","单价","销售金额","支付方式","会员ID"]
    rows = []
    for d_off in range(7):
        d = date(2026, 8, 1) + timedelta(days=d_off)
        for _ in range(random.randint(20, 40)):
            store = random.choice(STORES); cat = random.choice(list(CATEGORIES)); sub = random.choice(CATEGORIES[cat])
            names = SKU_NAMES[sub]; name = random.choice(names)
            qty = random.randint(1, 10); price = _price(sub)
            rows.append([str(d), store, _sku_id(cat, sub, names.index(name)), name, cat, sub, qty, price, round(qty*price,2),
                        random.choice(["微信","支付宝","银行卡","现金"]), f"M{random.randint(10000,99999)}" if random.random()<0.5 else ""])
    ws1 = wb.active; ws1.title = "销售明细"; xl_style(ws1, hd, rows)
    # returns
    ret_hd = ["退货日期","门店","SKU编号","商品名称","退货数量","退货金额","退货原因","原销售日期"]
    ret_rows = []
    for d_off in range(7):
        d = date(2026, 8, 1) + timedelta(days=d_off)
        for _ in range(random.randint(1, 4)):
            sub = random.choice(CATEGORIES["生鲜"]); name = random.choice(SKU_NAMES[sub])
            qty = random.randint(1, 2); price = _price(sub)
            ret_rows.append([str(d), random.choice(STORES), _sku_id("生鲜", sub, 0), name, qty, round(qty*price,2),
                           random.choice(["质量问题","顾客不满意"]), str(d - timedelta(days=random.randint(1,3)))])
    ws2 = wb.create_sheet("退货"); xl_style(ws2, ret_hd, ret_rows)
    # members
    mem_hd = ["会员ID","姓名","性别","手机号","注册日期","会员等级","累计积分","本月消费金额"]
    mem_rows = []
    for i in range(50):
        mid = f"M{20000+i}"; name = f"测试顾客{random.choice('赵钱孙李')}{random.choice('明华')}"
        mem_rows.append([mid, name, random.choice(["男","女"]), f"1{random.randint(30,99)}{random.randint(10000000,99999999)}",
                        str(date(2026, random.randint(1,8), random.randint(1,28))),
                        random.choices(["普通","银卡","金卡","钻石"],weights=[50,30,15,5])[0],
                        random.randint(0,10000), round(random.uniform(0,2000),2)])
    ws3 = wb.create_sheet("会员"); xl_style(ws3, mem_hd, mem_rows)
    wb.save(DATA_DIR / "a1" / "门店销售_202608.xlsx")
    print(f"  ✓ A1 holdout: 门店销售_202608.xlsx ({len(rows)} rows)")

    # edge: missing sheets
    wb2 = openpyxl.Workbook(); ws = wb2.active; ws.title = "销售明细"; xl_style(ws, hd, rows[:10])
    wb2.save(DATA_DIR / "a1" / "门店销售_仅销售sheet.xlsx")

    # edge: empty returns sheet
    wb3 = openpyxl.Workbook()
    ws_s = wb3.active; ws_s.title = "销售明细"; xl_style(ws_s, hd, rows)
    ws_r = wb3.create_sheet("退货"); xl_style(ws_r, ret_hd, [])
    ws_m = wb3.create_sheet("会员"); xl_style(ws_m, mem_hd, mem_rows)
    wb3.save(DATA_DIR / "a1" / "门店销售_空退货sheet.xlsx")
    print(f"  ✓ A1 edges: 仅销售sheet + 空退货sheet")

# ════════════════════════════════════════════════════════════
# A3
# ════════════════════════════════════════════════════════════
def gen_a3():
    hd = ["单据日期","单据编号","门店","商品编码","商品名称","销售数量","销售金额","收款方式","经手人"]
    erp = []
    for d_off in range(7):
        d = date(2026, 8, 1) + timedelta(days=d_off)
        for i in range(random.randint(15, 25)):
            sub = random.choice(CATEGORIES["食品"]); name = random.choice(SKU_NAMES[sub])
            qty = random.randint(1, 8); amt = round(qty * _price(sub), 2)
            erp.append([str(d), f"JD-{d.strftime('%Y%m%d')}-{i:04d}", random.choice(STORES),
                       _sku_id("食品", sub, 0), name, qty, amt, random.choice(["现金","微信"]), random.choice(["张丽","王强"])])
    write_csv("a3", "金蝶_导出_202608.csv", hd, erp)

    pos_hd = ["交易时间","交易号","店铺","商品条码","品名","数量","实收金额","支付方式","收银员"]
    pos = []
    for row in erp:
        pos.append([row[0]+" 14:30:00", row[1], row[2], row[3], row[4], row[5],
                   round(row[6]*random.uniform(0.97,1.01),2), row[7], random.choice(["小刘","小赵"])])
    for _ in range(10):
        d = date(2026, 8, 1) + timedelta(days=random.randint(0,6))
        sub = random.choice(CATEGORIES["食品"]); name = random.choice(SKU_NAMES[sub])
        pos.append([str(d)+" 19:00:00", f"POS-{random.randint(20000,29999)}", random.choice(STORES),
                   _sku_id("食品", sub, 0), name, random.randint(1,3), round(random.uniform(5,30),2), "微信", "小刘"])
    write_csv("a3", "pos_sales_202608.csv", pos_hd, pos)
    print(f"  ✓ A3 holdout: ERP {len(erp)} rows / POS {len(pos)} rows")

    # edges
    write_csv("a3", "金蝶_缺列.csv", hd[:6], [r[:6] for r in erp[:5]])
    mixed = []; 
    for row in erp[:20]:
        r = list(row)
        if random.random() < 0.5: r[0] = r[0].replace("-", "/")
        mixed.append(r)
    write_csv("a3", "金蝶_混合日期格式.csv", hd, mixed)
    print(f"  ✓ A3 edges: 缺列 + 混合日期格式")

# ════════════════════════════════════════════════════════════
# B4
# ════════════════════════════════════════════════════════════
def gen_b4():
    hd = ["活动编号","活动名称","开始日期","结束日期","覆盖门店","适用品类","促销类型","折扣力度","营销费用","活动前日均销售额","活动中日均销售额","活动后一周日均销售额"]
    rows = []
    for i in range(1, 6):
        start = date(2026, 8, 1) + timedelta(days=random.randint(0,20)); dur = random.choice([3,5,7]); end = start + timedelta(days=dur)
        base = random.uniform(2000, 12000); lift = random.uniform(0.05, 0.6); post = random.uniform(-0.15, 0.2)
        rows.append([f"P2026{i+20:04d}", random.choice(["开学季大促","中秋预热","品牌周"]), str(start), str(end),
                    ",".join(random.sample(STORES, k=2)), random.choice(["全品类","食品","日用品"]),
                    random.choice(["满减","折扣","秒杀"]), round(random.uniform(0.1,0.3),2),
                    round(random.uniform(1000,20000),2), round(base,2), round(base*(1+lift),2), round(base*(1+post),2)])
    rows.append(["P20269999","无效促销","2026-08-10","2026-08-12",STORES[0],"全品类","折扣",0.3,5000,10000.0,10000.0,10000.0])
    rows.append(["P20269998","亏本促销","2026-08-15","2026-08-17",STORES[0],"全品类","满减",0.4,30000,8000.0,9000.0,6000.0])
    write_csv("b4", "promo_events_august.csv", hd, rows)
    print(f"  ✓ B4 holdout: {len(rows)} events (含零提升+负ROI边界)")

# ════════════════════════════════════════════════════════════
# B5
# ════════════════════════════════════════════════════════════
def gen_b5():
    hd = ["SKU编号","商品名称","品类","子品类","供应商","进货价","零售价","季度销量","季度销售额","毛利率","库存周转天数","退货率","是否季节性商品","上架日期"]
    rows = []; idx = 1000
    for cat, subs in CATEGORIES.items():
        for sub in subs:
            for name in SKU_NAMES[sub][:2]:
                idx += 1; price = _price(sub); cost = round(price*random.uniform(0.50,0.75),2)
                qty = random.randint(10, 300); margin = round((price-cost)/price, 4)
                rows.append([_sku_id(cat, sub, idx), name, cat, sub, f"供应商{random.choice('ABCD')}", cost, price,
                           qty, round(price*qty,2), margin, round(random.uniform(3,150),1),
                           round(random.uniform(0,0.1),4), "是" if random.random()<0.3 else "否",
                           str(date(2025, random.randint(6,12),1))])
    rows.append(["SP999999","零销试销品","食品","零食","供应商A",5.0,15.0,0,0,0.6667,999,0,"否","2026-08-01"])
    rows.append(["SP999998","高退货品","食品","零食","供应商A",5.0,20.0,100,2000,0.75,3,1.0,"否","2026-06-01"])
    write_csv("b5", "product_performance_august.csv", hd, rows)
    print(f"  ✓ B5 holdout: {len(rows)} SKUs (含零销量+100%退货率)")

# ════════════════════════════════════════════════════════════
# B6
# ════════════════════════════════════════════════════════════
def gen_b6():
    hd = ["交易ID","会员ID","交易日期","交易金额","交易门店","购买品类数"]
    rows = []; mids = [f"M{30000+i}" for i in range(100)]
    for i in range(500):
        mid = random.choice(mids); d = date(2026, 6, 1) + timedelta(days=random.randint(0,60))
        rows.append([f"T{10000+i:06d}", mid, str(d), min(round(random.expovariate(1/150)*150+10,2),4000),
                    random.choice(STORES), random.randint(1,4)])
    for i in range(20):
        rows.append([f"T{20000+i:06d}", f"M{40000+i}", "2026-08-01", round(random.uniform(20,200),2), STORES[0], 1])
    write_csv("b6", "transactions_august.csv", hd, rows)
    print(f"  ✓ B6 holdout: {len(rows)} txns (含20个单次购买)")

# ════════════════════════════════════════════════════════════
# C7
# ════════════════════════════════════════════════════════════
def gen_c7():
    hd = ["月份","门店","目标销售额","实际销售额","达成率"]
    rows = []
    for store in STORES:
        target = round(random.uniform(150000,700000),2); actual = round(target*random.uniform(0.7,1.3),2)
        rows.append(["2026-08", store, target, actual, round(actual/target,4)])
    write_csv("c7", "sales_target_august.csv", hd, rows)
    write_csv("c7", "sales_target_缺门店.csv", hd, [["2026-08","",500000,480000,0.96]])
    print(f"  ✓ C7 holdout: sales_target_august.csv + edge (缺门店)")

# ════════════════════════════════════════════════════════════
# C8
# ════════════════════════════════════════════════════════════
def gen_c8():
    hd = ["日期","时间段","客流量","是否为节假日","天气"]
    rows = []; sun_rows = []
    for d_off in range(7):
        d = date(2026, 8, 3) + timedelta(days=d_off); ds = str(d)
        is_hol = "是" if d_off >= 5 else "否"; base = 70 if is_hol=="是" else 35
        for hour in [10,11,12,13,14,15,16,17,18,19,20,21]:
            peak = 1.0
            if hour in [11,12,18,19]: peak = 2.2 if is_hol=="是" else 1.7
            elif hour in [17,20]: peak = 1.8 if is_hol=="是" else 1.4
            r = [ds, f"{hour}:00-{hour+1}:00", int(base*peak*random.uniform(0.7,1.3)), is_hol, random.choice(["晴","多云","阴"])]
            rows.append(r)
            if is_hol == "是": sun_rows.append(r)
    write_csv("c8", "foot_traffic_august.csv", hd, rows)
    write_csv("c8", "foot_traffic_仅周末.csv", hd, sun_rows)
    print(f"  ✓ C8 holdout: foot_traffic_august.csv + edge (仅周末)")

# ════════════════════════════════════════════════════════════
# D9
# ════════════════════════════════════════════════════════════
def gen_d9():
    hd = ["SKU编号","商品名称","品类","子品类","门店","月初库存","月末库存","月销量","月销售额","上架天数"]
    rows = []; idx = 5000
    for cat, subs in CATEGORIES.items():
        for sub in subs:
            for name in SKU_NAMES[sub][:1]:
                for store in random.sample(STORES, k=2):
                    idx += 1; begin = random.randint(5,150); sold = random.randint(0,begin)
                    end = max(0, begin-sold+random.randint(0,15))
                    rows.append([_sku_id(cat, sub, idx), name, cat, sub, store, begin, end, sold, round(sold*_price(sub),2), random.randint(30,300)])
    rows.append(["SX999999","零动销品","生鲜","蔬菜",STORES[0],50,50,0,0,90])
    rows.append(["SX999998","负库存品","食品","零食",STORES[0],20,-5,30,450,60])
    write_csv("d9", "sku_sales_august.csv", hd, rows)
    print(f"  ✓ D9 holdout: {len(rows)} combinations (含零动销+负库存)")

# ════════════════════════════════════════════════════════════
# E11
# ════════════════════════════════════════════════════════════
def gen_e11():
    hd = ["员工编号","姓名","门店","岗位","基本工资","入职日期","社保基数"]
    roles = ["店长","副店长","收银员","理货员","生鲜技工","促销员","保洁"]
    bp = {"店长":8000,"副店长":6000,"收银员":4500,"理货员":4200,"生鲜技工":5000,"促销员":4000,"保洁":3500}
    rows = []
    for i in range(20):
        role = random.choice(roles); base = bp[role]
        rows.append([f"E{9000+i:04d}", f"测试员工{random.choice('赵钱孙李')}{random.choice('明华芳伟')}", random.choice(STORES),
                    role, base, str(date(2022,1,1)+timedelta(days=random.randint(0,1500))), round(base*random.uniform(0.8,1.2),2)])
    rows.append(["E99999","零薪员工",STORES[0],"保洁",0,"2026-08-01",0])
    write_csv("e11", "employees_august.csv", hd, rows)

    att_hd = ["员工编号","日期","上班时间","下班时间","是否迟到","是否早退","请假类型"]
    att_rows = []
    for emp in rows[:5]:
        for d_off in range(5):
            d = date(2026, 8, 3) + timedelta(days=d_off)
            late = random.random() < 0.1; early = random.random() < 0.05
            leave = random.choice(["无","无","无","无","事假","病假"])
            att_rows.append([emp[0], str(d), "09:30" if not late else f"09:{random.randint(35,55):02d}",
                           "18:00" if not early else f"17:{random.randint(30,55):02d}",
                           "是" if late else "否", "是" if early else "否", leave])
    write_csv("e11", "attendance_august.csv", att_hd, att_rows)
    print(f"  ✓ E11 holdout: {len(rows)} employees + {len(att_rows)} attendance records")

# ════════════════════════════════════════════════════════════
# E12
# ════════════════════════════════════════════════════════════
def gen_e12():
    hd1 = ["日期","门店","客流量","交易笔数","客单价","销售额","毛利率","同比销售额","环比销售额"]
    rows1 = []
    for d_off in range(7):
        d = date(2026, 8, 3) + timedelta(days=d_off)
        for store in STORES[:4]:
            traffic = random.randint(150,1200); txns = random.randint(int(traffic*0.3), int(traffic*0.6))
            asp = round(random.uniform(30,100),2); sales = round(txns*asp,2)
            rows1.append([str(d), store, traffic, txns, asp, sales, round(random.uniform(0.15,0.4),4),
                         round(sales*random.uniform(0.85,1.2),2), round(sales*random.uniform(0.9,1.15),2)])
    write_csv("e12", "store_sales_august.csv", hd1, rows1)

    hd2 = ["日期","门店","评价维度","评分","评价内容"]
    rows2 = []
    for d_off in range(7):
        d = date(2026, 8, 3) + timedelta(days=d_off)
        for store in STORES[:4]:
            for asp in random.sample(["商品质量","价格","服务态度","环境卫生","排队时间"], k=2):
                s = random.randint(3,5); rows2.append([str(d), store, asp, s, {5:"非常满意",4:"比较满意",3:"一般"}[s]])
    write_csv("e12", "customer_feedback_august.csv", hd2, rows2)

    hd3 = ["门店","品类","SKU数","库存数量","库存金额","低库存SKU数","过期预警SKU数"]
    rows3 = []
    for store in STORES[:4]:
        for cat in CATEGORIES:
            rows3.append([store, cat, random.randint(15,120), random.randint(150,4000),
                        round(random.uniform(3000,150000),2), random.randint(0,10), random.randint(0,5)])
    write_csv("e12", "inventory_snapshot_august.csv", hd3, rows3)

    # edge: BOM + blank lines
    p = DATA_DIR / "e12" / "store_sales_带BOM和空行.csv"
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(hd1); w.writerow([])
        for row in rows1[:10]: w.writerow(row)
        w.writerow([])
    print(f"  ✓ E12 holdout: 3 files + edge (BOM/空行)")

# ════════════════════════════════════════════════════════════
if __name__ == "__main__":
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    for sub in ["a1","a3","b4","b5","b6","c7","c8","d9","e11","e12"]:
        (DATA_DIR/sub).mkdir(parents=True, exist_ok=True)
    print("Generating HOLDOUT test data (separate from training data)...\n")
    gen_a1(); gen_a3(); gen_b4(); gen_b5(); gen_b6(); gen_c7(); gen_c8(); gen_d9(); gen_e11(); gen_e12()
    print("\nDone. Holdout data in tests/retail_skills/data_test/")
